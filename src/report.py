from datetime import datetime
from io import BytesIO
from numbers import Number

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font

from analytics import (
    calculate_category_totals,
    calculate_month_category_totals,
    calculate_month_totals,
    calculate_summary,
)


MONEY_FORMAT = "0.00"


def format_workbook(workbook, format_numbers=False):
    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for column in worksheet.columns:
            column_letter = column[0].column_letter
            longest_value = max(
                (len(str(cell.value)) for cell in column if cell.value is not None),
                default=0,
            )
            worksheet.column_dimensions[column_letter].width = longest_value + 2
            if format_numbers:
                for cell in column:
                    if isinstance(cell.value, Number):
                        cell.number_format = MONEY_FORMAT


def add_category_chart(workbook):
    worksheet = workbook["By Category"]
    chart = BarChart()
    chart.title = "Expenses by Category"
    chart.y_axis.title = "Amount"
    chart.x_axis.title = "Category"
    data = Reference(worksheet, min_col=2, min_row=1, max_row=worksheet.max_row)
    categories = Reference(worksheet, min_col=1, min_row=2, max_row=worksheet.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 14
    worksheet.add_chart(chart, "D2")


def create_excel_report(expenses, output_file):
    total_amount, average_expense, biggest_expense = calculate_summary(expenses)
    category_totals = calculate_category_totals(expenses)
    month_totals = calculate_month_totals(expenses)
    month_category_totals = calculate_month_category_totals(expenses).reset_index()
    category_difference = total_amount - category_totals["amount"].sum()
    month_difference = total_amount - month_totals["amount"].sum()
    row_count = len(expenses)
    month_count = expenses["month"].nunique()

    summary = pd.DataFrame(
        [
            {"metric": "Report created at", "value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {"metric": "Total amount", "value": total_amount},
            {"metric": "Average expense", "value": average_expense},
            {"metric": "Biggest expense date", "value": biggest_expense["date"]},
            {"metric": "Biggest expense category", "value": biggest_expense["category"]},
            {"metric": "Biggest expense description", "value": biggest_expense["description"]},
            {"metric": "Biggest expense amount", "value": biggest_expense["amount"]},
        ]
    )
    validation = pd.DataFrame(
        [
            {"check": "Raw Data sum equals By Category sum", "value": category_difference, "status": "OK" if abs(category_difference) < 0.01 else "ERROR"},
            {"check": "Raw Data sum equals By Month sum", "value": month_difference, "status": "OK" if abs(month_difference) < 0.01 else "ERROR"},
            {"check": "Row count is greater than zero", "value": row_count, "status": "OK" if row_count > 0 else "ERROR"},
            {"check": "Month count is greater than zero", "value": month_count, "status": "OK" if month_count > 0 else "ERROR"},
        ]
    )
    output_file.parent.mkdir(exist_ok=True)
    try:
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            expenses.to_excel(writer, sheet_name="Raw Data", index=False)
            summary.to_excel(writer, sheet_name="Summary", index=False)
            category_totals.to_excel(writer, sheet_name="By Category", index=False)
            month_totals.to_excel(writer, sheet_name="By Month", index=False)
            month_category_totals.to_excel(writer, sheet_name="Month x Category", index=False)
            validation.to_excel(writer, sheet_name="Validation", index=False)
            format_workbook(writer.book, format_numbers=True)
            add_category_chart(writer.book)
    except PermissionError as error:
        raise PermissionError(f"Close {output_file} in Excel and run the script again.") from error
    if not output_file.exists():
        raise FileNotFoundError(f"Excel report was not created: {output_file}")
    return total_amount, average_expense, biggest_expense, category_totals


def build_filtered_excel_report(expenses, filters):
    if expenses.empty:
        raise ValueError("Cannot create a report without filtered expenses.")

    category_totals = calculate_category_totals(expenses)
    month_totals = calculate_month_totals(expenses)
    month_category_totals = calculate_month_category_totals(expenses)
    summary = pd.DataFrame(
        [
            {"metric": "Report created at", "value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {"metric": "Selected date range", "value": filters["date_range"]},
            {"metric": "Selected categories", "value": filters["categories"]},
            {"metric": "Selected imports", "value": filters["imports"]},
            {"metric": "Text filter", "value": filters["description_text"]},
            {"metric": "Row count", "value": len(expenses)},
            {"metric": "Total amount", "value": expenses["amount"].sum()},
        ]
    )
    raw_total = expenses["amount"].sum()
    category_difference = raw_total - category_totals["amount"].sum()
    month_difference = raw_total - month_totals["amount"].sum()
    validation = pd.DataFrame(
        [
            {"check": "Filtered Data sum equals By Category sum", "value": category_difference, "status": "OK" if abs(category_difference) < 0.01 else "ERROR"},
            {"check": "Filtered Data sum equals By Month sum", "value": month_difference, "status": "OK" if abs(month_difference) < 0.01 else "ERROR"},
            {"check": "Row count is greater than zero", "value": len(expenses), "status": "OK"},
        ]
    )
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        expenses.to_excel(writer, sheet_name="Filtered Data", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False)
        category_totals.to_excel(writer, sheet_name="By Category", index=False)
        month_totals.to_excel(writer, sheet_name="By Month", index=False)
        month_category_totals.to_excel(writer, sheet_name="Month x Category", index=True)
        validation.to_excel(writer, sheet_name="Validation", index=False)
        format_workbook(writer.book)
    return output.getvalue()
